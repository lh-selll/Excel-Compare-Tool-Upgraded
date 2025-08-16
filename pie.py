from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.series import DataPoint

# 创建工作簿和工作表
wb = Workbook()
ws = wb.active

# 示例数据
data = [
    ["产品", "销量"],
    ["产品A", 150],
    ["产品B", 200],
    ["产品C", 180],
    ["产品D", 220],
    ["产品E", 170]
]

for row in data:
    ws.append(row)

# 计算数据点数量（排除表头）
data_points_count = len(data) - 1

# 创建柱状图
chart = BarChart()
chart.title = "不同类别柱子填充颜色示例"

# 设置数据范围
data_range = Reference(ws, min_col=2, min_row=1, max_col=2, max_row=6)
categories = Reference(ws, min_col=1, min_row=2, max_row=6)
chart.add_data(data_range, titles_from_data=True)
chart.set_categories(categories)

# 定义每个类别的颜色（RGB十六进制字符串格式）
colors = [
    "FF6384",  # 粉红色
    "36A2EB",  # 蓝色
    "FFCE56",  # 黄色
    "4BC0C0",  # 青色
    "9966FF"   # 紫色
]

# 为每个类别的柱子设置颜色
if chart.series:
    series = chart.series[0]
    # 为每个数据点创建DataPoint并应用颜色
    for i in range(min(data_points_count, len(colors))):
        # 仅需指定索引即可创建DataPoint
        dp = DataPoint(idx=i)
        series.dPt.append(dp)
        # 创建图形属性并设置填充颜色
        gp = GraphicalProperties(solidFill=colors[i])
        # 为数据点应用颜色
        series.dPt[i].graphicalProperties = gp

# 添加数据标签
chart.dataLabels = DataLabelList(showVal=True)

# 其他布局设置
chart.width = 20
chart.height = 15
chart.x_axis.title = "产品类别"
chart.y_axis.title = "销量"

# 添加图表到工作表
ws.add_chart(chart, "D2")

# 保存文件
wb.save("带颜色的柱状图.xlsx")
print("已生成每个类别带不同颜色的柱状图：带颜色的柱状图.xlsx")
    